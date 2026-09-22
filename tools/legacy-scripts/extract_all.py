# -*- coding: utf-8 -*-
"""把 5 份源头资料完整抽出：文字 → txt，图片 → assets/。
只打印紧凑清单（避免刷屏），详细内容留在文件里供逐步阅读。"""
import base64
import io
import os
import re
import zipfile

OUT = r"E:\HikRobot\_萃取\newdoc"
AST = os.path.join(OUT, "assets")
os.makedirs(AST, exist_ok=True)

SRC = {
    "pdf":  r"E:\HikRobot\Vision\文档资料\产品手册\机器视觉产品手册合集\机器视觉工控机产品手册-V.151.CN.25Q2.1(阅读版).pdf",
    "pptx": r"E:\HikRobot\AGV\解决方案\01 LMR-XXXX搬运项目方案V1.4.pptx",
    "v32":  r"D:\Desktop\誉创\金卫项目\方案\江苏誉创_金卫智慧舱AGV物料配送与PLC通讯调度方案_V3.2.docx",
    "net":  r"D:\Desktop\誉创\金卫项目\方案\方案文件\江苏誉创_金卫智慧舱_网络解决方案.html",
    "agv":  r"D:\Desktop\誉创\金卫项目\方案\方案文件\江苏誉创_金卫智慧舱_AGV方案_V5.0.html",
}
inv = []


def log(s):
    inv.append(s)
    print(s)


print("=== 源文件存在性 ===")
for k, p in SRC.items():
    ok = os.path.exists(p)
    print("  %-6s %s  %s" % (k, "✓" if ok else "✗ 缺失", ("%.1f MB" % (os.path.getsize(p)/1048576)) if ok else p))

# ── 1) PDF ──────────────────────────────────────────────────────────
try:
    import pymupdf
    d = pymupdf.open(SRC["pdf"])
    txt = []
    nimg = 0
    for i, pg in enumerate(d):
        txt.append("\n===== 第 %d 页 =====\n" % (i + 1) + (pg.get_text() or ""))
        for j, im in enumerate(pg.get_images(full=True)):
            try:
                info = d.extract_image(im[0])
                fn = "pdf_p%03d_%d.%s" % (i + 1, j + 1, info["ext"])
                open(os.path.join(AST, fn), "wb").write(info["image"])
                nimg += 1
            except Exception:
                pass
    io.open(os.path.join(OUT, "src_pdf.txt"), "w", encoding="utf-8").write("".join(txt))
    log("  [PDF] 页数=%d  抽出图片=%d  文本=%d 字符 → src_pdf.txt" % (d.page_count, nimg, sum(len(x) for x in txt)))
except Exception as e:
    log("  [PDF] 失败：%s" % e)

# ── 2) PPTX ─────────────────────────────────────────────────────────
try:
    z = zipfile.ZipFile(SRC["pptx"])
    slides = sorted([n for n in z.namelist() if re.match(r"ppt/slides/slide\d+\.xml$", n)],
                    key=lambda x: int(re.findall(r"\d+", x)[-1]))
    out = []
    for n in slides:
        xml = z.read(n).decode("utf-8", "ignore")
        tx = re.findall(r"<a:t>(.*?)</a:t>", xml, re.S)
        out.append("\n===== %s =====\n" % n.split("/")[-1] + "\n".join(tx))
    io.open(os.path.join(OUT, "src_pptx.txt"), "w", encoding="utf-8").write("\n".join(out))
    media = [n for n in z.namelist() if n.startswith("ppt/media/")]
    for n in media:
        open(os.path.join(AST, "pptx_" + os.path.basename(n)), "wb").write(z.read(n))
    log("  [PPTX] 幻灯片=%d  文本=%d 字符  媒体=%d 个 → src_pptx.txt" % (len(slides), sum(len(x) for x in out), len(media)))
except Exception as e:
    log("  [PPTX] 失败：%s" % e)

# ── 3) DOCX ─────────────────────────────────────────────────────────
try:
    import docx
    d = docx.Document(SRC["v32"])
    lines = [p.text for p in d.paragraphs]
    for t in d.tables:
        lines.append("\n[表格]")
        for r in t.rows:
            lines.append(" | ".join(c.text.strip().replace("\n", " ") for c in r.cells))
    io.open(os.path.join(OUT, "src_v32.txt"), "w", encoding="utf-8").write("\n".join(lines))
    z = zipfile.ZipFile(SRC["v32"])
    media = [n for n in z.namelist() if n.startswith("word/media/")]
    for n in media:
        open(os.path.join(AST, "v32_" + os.path.basename(n)), "wb").write(z.read(n))
    log("  [DOCX] 段落=%d  表格=%d  媒体=%d → src_v32.txt" % (len(d.paragraphs), len(d.tables), len(media)))
except Exception as e:
    log("  [DOCX] 失败：%s" % e)

# ── 4/5) HTML（两个）────────────────────────────────────────────────
for key, path, outname in (("net", SRC["net"], "src_net.txt"), ("agv", SRC["agv"], "src_agv.txt")):
    try:
        h = io.open(path, encoding="utf-8").read()
        # 正文文本（按 .sheet 分页）
        body = re.sub(r"<style.*?</style>", "", h, flags=re.S)
        body = re.sub(r"<script.*?</script>", "", body, flags=re.S)
        sheets = re.findall(r'<div class="sheet">(.*?)(?=<div class="sheet">|</body>)', body, re.S)
        txt = []
        for i, sh in enumerate(sheets, 1):
            t = re.sub(r"<[^>]+>", "\n", sh)
            t = re.sub(r"\n{2,}", "\n", t).strip()
            txt.append("\n===== 第 %d 页 =====\n" % i + t)
        io.open(os.path.join(OUT, outname), "w", encoding="utf-8").write("\n".join(txt))
        # base64 图落地
        n = 0
        for m in re.finditer(r'data:image/(\w+);base64,([A-Za-z0-9+/=]+)', h):
            n += 1
            open(os.path.join(AST, "%s_img%d.%s" % (key, n, m.group(1))), "wb").write(base64.b64decode(m.group(2)))
        log("  [%s] sheet=%d  表=%d  内联base64图=%d → %s" % (key, len(sheets), h.count("<table"), n, outname))
    except Exception as e:
        log("  [%s] 失败：%s" % (key, e))

# ── 资产清单 ────────────────────────────────────────────────────────
files = sorted(os.listdir(AST))
tot = sum(os.path.getsize(os.path.join(AST, f)) for f in files)
log("\n=== 资源库 assets/：%d 个文件，%.1f MB ===" % (len(files), tot / 1048576))
from collections import Counter
c = Counter(f.split("_")[0].split(".")[0] for f in files)
log("  按来源：" + "  ".join("%s=%d" % kv for kv in sorted(c.items())))
big = sorted(files, key=lambda f: -os.path.getsize(os.path.join(AST, f)))[:12]
log("  最大的 12 个：" + ", ".join("%s(%.0fKB)" % (f, os.path.getsize(os.path.join(AST, f)) / 1024) for f in big))
io.open(os.path.join(OUT, "inventory.txt"), "w", encoding="utf-8").write("\n".join(inv))
